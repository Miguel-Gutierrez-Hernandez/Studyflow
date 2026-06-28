"""
ingesta/extractor.py — Extracción de texto de múltiples formatos.

Formatos soportados:
    .pdf    → PyMuPDF (fitz)
    .docx   → python-docx
    .pptx   → python-pptx
    .txt    → lectura directa
    .mp3 / .wav / .m4a / .ogg / .flac → Whisper (transcripción en español)

Uso:
    from ingesta.extractor import extraer_texto

    texto = extraer_texto(Path("apuntes.pdf"))
    texto = extraer_texto(Path("clase.mp3"))
"""

from __future__ import annotations

from pathlib import Path

# Extensiones soportadas
FORMATOS_DOCUMENTO = {".pdf", ".docx", ".pptx", ".txt"}
FORMATOS_AUDIO = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".webm", ".opus"}


def extraer_texto(ruta: Path) -> str:
    """
    Punto de entrada principal. Detecta el formato y extrae el texto.

    Args:
        ruta: Path al archivo a procesar.

    Returns:
        Texto extraído como string limpio.

    Raises:
        ValueError: Si el formato no está soportado.
        FileNotFoundError: Si el archivo no existe.
    """
    ruta = Path(ruta)
    if not ruta.exists():
        raise FileNotFoundError(f"No se encuentra el archivo: {ruta}")

    sufijo = ruta.suffix.lower()

    if sufijo == ".pdf":
        return _extraer_pdf(ruta)
    elif sufijo == ".docx":
        return _extraer_docx(ruta)
    elif sufijo == ".pptx":
        return _extraer_pptx(ruta)
    elif sufijo == ".txt":
        return _extraer_txt(ruta)
    elif sufijo in FORMATOS_AUDIO:
        return _transcribir_audio(ruta)
    else:
        raise ValueError(
            f"Formato '{sufijo}' no soportado. "
            f"Documentos: {FORMATOS_DOCUMENTO} | Audio: {FORMATOS_AUDIO}"
        )


# ── PDF ───────────────────────────────────────────────────────────────────────

def _extraer_pdf(ruta: Path) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise ImportError("Instala PyMuPDF: pip install pymupdf")

    doc = fitz.open(str(ruta))
    paginas = []
    for num, pagina in enumerate(doc, start=1):
        texto = pagina.get_text("text").strip()
        if texto:
            paginas.append(f"[Página {num}]\n{texto}")
    doc.close()
    return "\n\n".join(paginas)


# ── DOCX ──────────────────────────────────────────────────────────────────────

def _extraer_docx(ruta: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        raise ImportError("Instala python-docx: pip install python-docx")

    doc = Document(str(ruta))
    parrafos = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(parrafos)


# ── PPTX ──────────────────────────────────────────────────────────────────────

def _extraer_pptx(ruta: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("Instala python-pptx: pip install python-pptx")

    prs = Presentation(str(ruta))
    diapositivas = []
    for num, slide in enumerate(prs.slides, start=1):
        textos = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                textos.append(shape.text.strip())
        if textos:
            diapositivas.append(f"[Diapositiva {num}]\n" + "\n".join(textos))
    return "\n\n".join(diapositivas)


# ── TXT ───────────────────────────────────────────────────────────────────────

def _extraer_txt(ruta: Path) -> str:
    # Intenta UTF-8, fallback a latin-1 para apuntes con caracteres españoles
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            return ruta.read_text(encoding=encoding).strip()
        except UnicodeDecodeError:
            continue
    raise ValueError(f"No se pudo leer {ruta} con ningún encoding conocido.")


# ── AUDIO → Whisper ───────────────────────────────────────────────────────────

def _transcribir_audio(ruta: Path) -> str:
    """
    Transcribe audio usando Whisper local.
    Fuerza idioma español para mayor precisión.
    """
    try:
        import whisper
    except ImportError:
        raise ImportError(
            "Instala Whisper: pip install openai-whisper\n"
            "También necesitas ffmpeg instalado en el sistema."
        )

    from config import WHISPER_MODEL

    print(f"  ⏳ Cargando modelo Whisper '{WHISPER_MODEL}'...")
    modelo = whisper.load_model(WHISPER_MODEL)

    print(f"  🎙️  Transcribiendo '{ruta.name}' (puede tardar unos minutos)...")
    resultado = modelo.transcribe(
        str(ruta),
        language="es",        # forzar español
        task="transcribe",    # no traducir, solo transcribir
        verbose=False,
    )
    return resultado["text"].strip()


# ── Batch ─────────────────────────────────────────────────────────────────────

def extraer_todos(rutas: list[Path]) -> dict[str, str]:
    """
    Extrae texto de una lista de archivos.

    Returns:
        Dict {nombre_archivo: texto_extraido}
        Los archivos con error se incluyen con el mensaje de error.
    """
    resultados = {}
    for ruta in rutas:
        try:
            print(f"  📄 Procesando: {ruta.name}")
            resultados[ruta.name] = extraer_texto(ruta)
            print(f"     ✅ OK ({len(resultados[ruta.name])} caracteres)")
        except Exception as e:
            print(f"     ❌ Error: {e}")
            resultados[ruta.name] = f"[ERROR: {e}]"
    return resultados
