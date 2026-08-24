"""
consumption/extractor.py — Text extraction from multiple file formats.

Supported:
    .pdf    -> PyMuPDF
    .docx   -> python-docx
    .pptx   -> python-pptx
    .txt    -> direct read
    audio   -> faster-whisper (local, offline transcription)

Audio transcription requires the 'faster-whisper' package. Install with:
    pip install faster-whisper --break-system-packages

The first transcription of a given model size downloads model weights from
Hugging Face and caches them locally; subsequent runs are offline.
"""

from pathlib import Path

DOCUMENT_FORMATS = {".pdf", ".docx", ".pptx", ".txt"}
AUDIO_FORMATS = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".webm", ".opus"}

# Cached whisper model instance, lazily created on first audio file.
_whisper_model = None
_WHISPER_MODEL_SIZE = "base"  # tiny | base | small | medium | large-v3


def _get_whisper_model():
    global _whisper_model
    if _whisper_model is None:
        try:
            from faster_whisper import WhisperModel
        except ImportError:
            raise ImportError(
                "Install faster-whisper for audio transcription: "
                "pip install faster-whisper --break-system-packages"
            )
        # compute_type="int8" keeps this usable on CPU-only machines (e.g. Apple Silicon).
        _whisper_model = WhisperModel(_WHISPER_MODEL_SIZE, device="cpu", compute_type="int8")
    return _whisper_model


def extract_text(path: Path) -> str:
    """Detect format and extract text."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf(path)
    elif ext == ".docx":
        return _docx(path)
    elif ext == ".pptx":
        return _pptx(path)
    elif ext == ".txt":
        return _txt(path)
    elif ext in AUDIO_FORMATS:
        return _audio(path)
    else:
        raise ValueError(
            f"Unsupported format '{ext}'. Supported: {DOCUMENT_FORMATS | AUDIO_FORMATS}"
        )


def extract_all(paths: list[Path]) -> dict[str, str]:
    """Extract text from a list of files. Returns {{filename: text}}."""
    results = {}
    for path in paths:
        try:
            print(f"  Processing: {path.name}")
            results[path.name] = extract_text(path)
            print(f"     OK ({len(results[path.name])} chars)")
        except Exception as e:
            print(f"     ERROR: {e}")
            results[path.name] = f"[ERROR: {e}]"
    return results


def _pdf(path: Path) -> str:
    try:
        import fitz
    except ImportError:
        raise ImportError("Install PyMuPDF: pip install pymupdf")
    doc = fitz.open(str(path))
    pages = [
        f"[Page {i+1}]\n{page.get_text('text').strip()}"
        for i, page in enumerate(doc)
        if page.get_text("text").strip()
    ]
    doc.close()
    return "\n\n".join(pages)


def _docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        raise ImportError("Install python-docx: pip install python-docx")
    doc = Document(str(path))
    return "\n\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())


def _pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("Install python-pptx: pip install python-pptx")
    prs = Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        if texts:
            slides.append(f"[Slide {i}]\n" + "\n".join(texts))
    return "\n\n".join(slides)


def _audio(path: Path) -> str:
    model = _get_whisper_model()
    segments, info = model.transcribe(str(path), beam_size=5)
    lines = [seg.text.strip() for seg in segments if seg.text.strip()]
    if not lines:
        return ""
    header = f"[Transcript · detected language: {info.language} ({info.language_probability:.0%})]"
    return header + "\n\n" + "\n".join(lines)


def _txt(path: Path) -> str:
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            return path.read_text(encoding=enc).strip()
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Could not read {path} with any known encoding.")